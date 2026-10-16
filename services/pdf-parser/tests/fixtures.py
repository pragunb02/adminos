"""SYNTHETIC statements modelled on common Indian bank export layouts (not real statements; layouts change)."""

import csv
import io

from fpdf import FPDF
from fpdf.enums import EncryptionMethod


def hdfc_csv() -> bytes:
    rows = [
        ["HDFC BANK Ltd."], ["Account No : XXXXXXXX1234", "", "IFSC : HDFC0001234"], [],
        ["Date", "Narration", "Chq./Ref.No.", "Value Dt", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"],
        ["01/10/26", "UPI-SWIGGY-swiggy@icici-ICIC0000001-123456789012-Order", "0000123456789012", "01/10/26", "249.00", "", "49,751.00"],
        ["", "food delivery", "", "", "", "", ""],  # continuation of the narration
        ["02/10/26", "NEFT CR-ACME CORP-SALARY OCT", "N123", "02/10/26", "", "1,00,000.00", "1,49,751.00"],
        ["03/10/26", "POS 1234XXXX AMAZON PAY IN", "", "03/10/26", "1,299.50", "", "1,48,451.50"],
        [], ["STATEMENT SUMMARY", "", "", "", "1,548.50", "1,00,000.00", ""],
    ]
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode()


def sbi_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [
        ["State Bank of India"], ["Account Number", "00000012345678"], ["Opening Balance", "10,000.00"], [],
        ["Txn Date", "Value Date", "Description", "Ref No./Cheque No.", "Debit", "Credit", "Balance"],
        ["1 Oct 2026", "1 Oct 2026", "TO TRANSFER-UPI/DR/123456789012/ZOMATO/YESB/zomato@ybl", "TRANSFER TO 456", "450.00", "", "9,550.00"],
        ["2 Oct 2026", "2 Oct 2026", "BY TRANSFER-NEFT*HDFC0000001*REFUND", "TRANSFER FROM 789", "", "450.00", "10,000.00"],
        ["3 Oct 2026", "3 Oct 2026", "ATM WDL-ATM CASH 1234 MUMBAI", "ATM", "2,000.00", "", "8,000.00"],
    ]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def icici_rows() -> list:
    return [
        ["DETAILED STATEMENT", "ICICI BANK LIMITED"], ["Account Number", "XXXXXXXX5678"], [],
        ["S No.", "Value Date", "Transaction Date", "Cheque Number", "Transaction Remarks", "Withdrawal Amount (INR )", "Deposit Amount (INR )", "Balance (INR )"],
        ["1", "05/10/2026", "05/10/2026", "-", "UPI/123456789012/NETFLIX/netflix@icici", "649.00", "0.00", "20,000.00"],
        ["2", "06/10/2026", "06/10/2026", "-", "INT CREDIT", "0.00", "12.34", "20,012.34"],
    ]


def kotak_drcr_csv() -> bytes:
    rows = [
        ["Kotak Mahindra Bank"], ["Account No. 1234567890"],
        ["Sl. No.", "Transaction Date", "Value Date", "Description", "Chq / Ref No.", "Amount", "Dr / Cr", "Balance", "Dr / Cr"],
        ["1", "07-10-2026", "07-10-2026", "UPI/UBER/uber@axis", "UPI-1", "320.00", "DR", "5,000.00", "CR"],
        ["2", "08-10-2026", "08-10-2026", "IMPS-RAVI KUMAR", "IMPS-2", "1,500.00", "CR", "6,500.00", "CR"],
        ["3", "09-10-2026", "09-10-2026", "MYSTERY", "X", "10.00", "", "6,490.00", "CR"],  # no direction → skipped
    ]
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode()


def axis_pdf(password: str | None = None) -> bytes:
    pdf = FPDF()
    if password:
        pdf.set_encryption(owner_password="owner-" + password, user_password=password, encryption_method=EncryptionMethod.AES_128)
    pdf.add_page()
    pdf.set_font("Helvetica", size=8)
    pdf.cell(0, 6, "AXIS BANK  Statement of Account No : 917010012349876", new_x="LMARGIN", new_y="NEXT")
    data = [
        ("Tran Date", "Particulars", "Debit", "Credit", "Balance"),
        ("10-10-2026", "UPI/P2M/123456789012/SWIGGY", "199.00", "", "4,801.00"),
        ("11-10-2026", "UPI/P2A/987654321098/MOM", "", "2,000.00", "6,801.00"),
        ("12-10-2026", "ECOM PUR/FLIPKART/BLR", "1,299.00", "", "5,502.00"),
    ]
    with pdf.table() as table:
        for r in data:
            row = table.row()
            for c in r:
                row.cell(c)
    return bytes(pdf.output())
