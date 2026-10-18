"""Hostile files and the parser's work limits (AOS-626, AUD-22). All files are synthetic."""

import csv
import io
import logging
import time

import pytest
from fastapi.testclient import TestClient
from fpdf import FPDF

from app import limits
from app.isolation import ParseCrashed, ParseFailed, ParseTimeout, run_isolated
from app.limits import Budget
from app.main import app
from app.models import ParseError
from app.readers import read, read_csv, read_pdf, read_xlsx
from tests import fixtures as fx
from tests import hostile_targets as hostile

client = TestClient(app)


def post(data: bytes, name: str):
    return client.post("/v1/statements:parse", files={"file": (name, data)})


def csv_bytes(rows) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode()


def blank_pdf(pages: int) -> bytes:
    pdf = FPDF()
    for _ in range(pages):
        pdf.add_page()
    return bytes(pdf.output())


# --- Row, cell, sheet and page caps ---

def test_budget_counts_rows_and_cells_and_refuses_past_either_cap():
    b = Budget(max_rows=2, max_cells=10)
    b.charge(4)
    b.charge(4)
    with pytest.raises(ParseError) as e:
        b.charge(1)
    assert e.value.code == "TOO_MANY_ROWS"
    with pytest.raises(ParseError):
        Budget(max_rows=100, max_cells=10).charge(11)


def test_a_sheet_that_declares_too_many_rows_or_cells_is_refused_before_it_is_read():
    with pytest.raises(ParseError) as e:
        Budget(max_rows=10, max_cells=1000).check_declared(rows=11, cols=1)
    assert e.value.code == "TOO_MANY_ROWS"
    with pytest.raises(ParseError):
        Budget(max_rows=1000, max_cells=100).check_declared(rows=10, cols=11)
    Budget(max_rows=10, max_cells=100).check_declared(rows=10, cols=10)  # exactly at the cap is fine


def test_a_csv_of_many_tiny_rows_is_refused_with_the_real_caps():
    data = b"1\n" * (limits.MAX_ROWS + 1)
    res = post(data, "s.csv")
    assert res.status_code == 422 and res.json()["error"]["code"] == "TOO_MANY_ROWS"


def test_a_csv_of_very_wide_rows_is_refused_by_the_cell_cap():
    with pytest.raises(ParseError) as e:
        read_csv(csv_bytes([["x"] * 1000] * 50), Budget(max_cells=10_000))
    assert e.value.code == "TOO_MANY_ROWS"


def test_rows_are_counted_across_all_sheets_of_a_workbook():
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["a"] * 5)
    for i in range(3):
        ws = wb.create_sheet(f"S{i}")
        for _ in range(4):
            ws.append(["b"] * 5)
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ParseError) as e:
        read_xlsx(buf.getvalue(), Budget(max_rows=12))
    assert e.value.code == "TOO_MANY_ROWS"
    assert len(read_xlsx(buf.getvalue(), Budget(max_rows=13)).tables) == 4


def test_a_workbook_with_too_many_sheets_is_refused(monkeypatch):
    import openpyxl

    monkeypatch.setattr(limits, "MAX_SHEETS", 3)
    wb = openpyxl.Workbook()
    for i in range(3):
        wb.create_sheet(f"S{i}")
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ParseError) as e:
        read_xlsx(buf.getvalue(), Budget())
    assert e.value.code == "TOO_MANY_SHEETS"


def test_a_pdf_with_too_many_pages_is_refused_with_the_real_cap():
    res = post(blank_pdf(limits.MAX_PAGES + 1), "s.pdf")
    assert res.status_code == 422 and res.json()["error"]["code"] == "TOO_MANY_PAGES"


def test_a_pdf_at_the_page_cap_is_read(monkeypatch):
    monkeypatch.setattr(limits, "MAX_PAGES", 2)
    assert read_pdf(blank_pdf(2), None, Budget()).tables == []
    with pytest.raises(ParseError) as e:
        read_pdf(blank_pdf(3), None, Budget())
    assert e.value.code == "TOO_MANY_PAGES"


def test_pdf_table_rows_are_charged_to_the_budget():
    with pytest.raises(ParseError) as e:
        read(fx.axis_pdf(), "a.pdf", None, Budget(max_rows=2))
    assert e.value.code == "TOO_MANY_ROWS"


def test_real_statements_stay_well_inside_the_caps():
    for data, name in [(fx.hdfc_csv(), "a.csv"), (fx.sbi_xlsx(), "a.xlsx"), (fx.axis_pdf(), "a.pdf"), (fx.kotak_drcr_csv(), "k.csv")]:
        assert post(data, name).status_code == 200, name


# --- Malformed files fail closed with a code, never content ---

@pytest.mark.parametrize("data,name", [
    (b"%PDF-1.7\n" + b"\x00\xff" * 500, "s.pdf"),  # a PDF header with garbage
    (b"PK\x03\x04" + b"\x00" * 200, "s.xlsx"),  # a truncated zip
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 200, "s.xls"),  # an OLE header with nothing behind it
])
def test_malformed_files_are_unsupported_not_errors(data, name):
    res = post(data, name)
    assert res.status_code == 422 and res.json()["error"]["code"] == "UNSUPPORTED_FORMAT"


def test_a_field_over_the_csv_field_limit_is_unreadable_and_not_echoed():
    data = b"Date,Narration\n01/10/26," + b"SWIGGY" * 30_000 + b"\n"
    res = post(data, "s.csv")
    assert res.status_code == 422 and "SWIGGY" not in res.text


# --- The time limit: a parse runs in a child process that is killed when the time is up ---

def test_a_parse_that_runs_too_long_is_killed():
    started = time.monotonic()
    with pytest.raises(ParseTimeout):
        run_isolated(hostile.spin_forever, (), timeout=1.0)
    assert time.monotonic() - started < 10


def test_the_real_endpoint_answers_parse_timeout_when_the_limit_passes(monkeypatch):
    monkeypatch.setattr(limits, "PARSE_TIMEOUT_SECONDS", 0.001)  # the child can't even start in time
    res = post(fx.hdfc_csv(), "s.csv")
    assert res.status_code == 422 and res.json()["error"]["code"] == "PARSE_TIMEOUT"


def test_a_child_that_dies_is_a_crash_not_a_hang():
    with pytest.raises(ParseCrashed):
        run_isolated(hostile.die, (), timeout=20)


def test_a_child_that_allocates_too_much_memory_fails():
    with pytest.raises(ParseFailed) as e:
        run_isolated(hostile.allocate, (2 * 1024 * 1024 * 1024,), timeout=20, memory_bytes=256 * 1024 * 1024)
    assert e.value.kind == "MemoryError"


def test_parse_errors_come_back_with_their_code_and_other_failures_with_their_type_only():
    with pytest.raises(ParseError) as e:
        run_isolated(hostile.refuse, (), timeout=20)
    assert e.value.code == "NO_TRANSACTIONS"
    with pytest.raises(ParseFailed) as f:
        run_isolated(hostile.fail_with_content, (), timeout=20)
    assert f.value.kind == "ValueError" and "SWIGGY" not in str(f.value)
    assert run_isolated(hostile.echo, ({"rows": 3},), timeout=20) == {"rows": 3}


def test_timeouts_and_crashes_are_logged_without_content(monkeypatch, caplog):
    import app.main as m

    caplog.set_level(logging.INFO)

    def crash(*_):
        raise ParseCrashed()

    monkeypatch.setattr(m, "run_isolated", crash)
    res = post(b"Date,Narration\n01/10/26,SWIGGY 249\n", "Ravi_Statement.csv")
    assert res.status_code == 422 and res.json()["error"]["code"] == "UNREADABLE"
    assert "SWIGGY" not in caplog.text and "Ravi_Statement" not in caplog.text
