"""File → raw tables (lists of string cells) plus header text, per format. No parsing of meaning here."""

import csv
import io
from dataclasses import dataclass, field
from typing import List

from app.models import ParseError

Table = List[List[str]]


@dataclass
class RawStatement:
    tables: List[Table]
    text: str = ""  # first page / header text, for bank and account detection
    warnings: List[str] = field(default_factory=list)


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def read_csv(data: bytes) -> RawStatement:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 never fails
        raise ParseError("UNSUPPORTED_FORMAT", "Unreadable text encoding")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[_clean(c) for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    return RawStatement(tables=[rows], text="\n".join(" ".join(r) for r in rows[:30]))


def read_xlsx(data: bytes) -> RawStatement:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ParseError("UNSUPPORTED_FORMAT", "Not a readable .xlsx file")
    tables = [[[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    return RawStatement(tables=tables, text="\n".join(" ".join(r) for t in tables for r in t[:30]))


def read_xls(data: bytes) -> RawStatement:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception:
        raise ParseError("UNSUPPORTED_FORMAT", "Not a readable .xls file")
    tables = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            cells = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    cells.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).date().isoformat())
                else:
                    cells.append(_clean(cell.value))
            rows.append(cells)
        tables.append(rows)
    return RawStatement(tables=tables, text="\n".join(" ".join(r) for t in tables for r in t[:30]))


def read_pdf(data: bytes, password: str | None) -> RawStatement:
    import pdfplumber

    try:
        pdf = pdfplumber.open(io.BytesIO(data), password=password or "")
        first = pdf.pages[0] if pdf.pages else None
    except Exception as e:
        if _is_password_error(e):
            if password:
                raise ParseError("PASSWORD_INCORRECT", "The statement password is incorrect")
            raise ParseError("PASSWORD_REQUIRED", "This statement is password-protected")
        raise ParseError("UNSUPPORTED_FORMAT", "Not a readable PDF")
    with pdf:
        tables: List[Table] = []
        for page in pdf.pages:
            for t in page.extract_tables() or []:
                tables.append([[_clean(c) for c in row] for row in t])
        text = (first.extract_text() or "") if first else ""
    warnings = [] if tables else ["No tables found in the PDF; it may be a scanned image"]
    return RawStatement(tables=tables, text=text, warnings=warnings)


def _is_password_error(e: BaseException) -> bool:
    """pdfplumber wraps pdfminer's PDFPasswordIncorrect; look through the wrapper chain."""
    from pdfminer.pdfdocument import PDFPasswordIncorrect

    seen = set()
    while e is not None and id(e) not in seen:
        seen.add(id(e))
        if isinstance(e, PDFPasswordIncorrect) or any(isinstance(a, PDFPasswordIncorrect) for a in getattr(e, "args", ())):
            return True
        e = e.__cause__ or e.__context__
    return False


def read(data: bytes, filename: str, password: str | None) -> RawStatement:
    name = (filename or "").lower()
    if data[:5] == b"%PDF-" or name.endswith(".pdf"):
        return read_pdf(data, password)
    if data[:4] == b"PK\x03\x04" or name.endswith(".xlsx"):
        return read_xlsx(data)
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or name.endswith(".xls"):
        return read_xls(data)
    if name.endswith((".csv", ".txt")) or b"\x00" not in data[:1024]:
        return read_csv(data)
    raise ParseError("UNSUPPORTED_FORMAT", "Upload a PDF, CSV, XLS or XLSX statement")
